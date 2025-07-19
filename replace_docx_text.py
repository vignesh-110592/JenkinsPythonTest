import sys
import os
from datetime import datetime
from docx import Document

def replace_placeholders(doc, replacements):
    for para in doc.paragraphs:
        for key, val in replacements.items():
            if key in para.text:
                inline = para.runs
                for i in range(len(inline)):
                    if key in inline[i].text:
                        print(f"Replacing '{key}' with '{val}'")
                        inline[i].text = inline[i].text.replace(key, val)

def process_file(input_filename, output_filename, replacements):
    doc = Document(input_filename)
    replace_placeholders(doc, replacements)
    doc.save(output_filename)

if __name__ == "__main__":
    if len(sys.argv) != 5:
        print("Usage: python replace_docx_text.py <CompanyName> <Greeting> <PersonName> <PositionName>")
        sys.exit(1)

    company_name = sys.argv[1]
    Greeting = sys.argv[2]
    person_name = sys.argv[3]
    position_name = sys.argv[4]

    current_date = datetime.today().strftime("%d.%m.%Y")

    replacements_anschreiben = {
        "__DATE__": current_date,
        "Company_name": company_name,
        "Greeting": Greeting,
        "Person_name": person_name,
        "Position_name": position_name
    }

    replacements_lebenslauf = {
        "__DATE__": current_date
    }

    base_path = os.getcwd()
    process_file(os.path.join(base_path, "AnschreibenRaw.docx"),
                 os.path.join(base_path, "Anschreiben.docx"),
                 replacements_anschreiben)
    
    # Process Lebenslauf if needed
    modified_timestamp = os.path.getmtime("Lebenslauf.docx")
    modified_date = datetime.fromtimestamp(modified_timestamp).strftime("%d.%m.%Y")
    print(modified_date)
    # Compare dates
    if modified_date == current_date:
        print(f"File 'Lebenslauf.docx' was modified today. Skipping Lebenslauf processing.")
    else:
        print(f"File 'Lebenslauf.docx' was last modified on {modified_date} and will be processed.")
        process_file(os.path.join(base_path, "LebenslaufRaw.docx"),
                 os.path.join(base_path, "Lebenslauf.docx"),
                 replacements_lebenslauf)

    

    print("Replacement complete.")
